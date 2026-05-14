"""
Router de entrada: validación de seguridad, segmentación por zonas y decisión
de ruta.

Para cada fichero subido:
    1. Valida tamaño, tipo MIME real por firma/contenido, no por extensión.
    2. Rechaza ficheros sospechosos o no admitidos.
    3. Extrae texto cuando hay capa textual útil.
    4. SEGMENTA POR ZONAS: en lugar de enviar el documento completo al modelo,
       localiza las zonas donde se concentran los datos descriptivos
       (preliminares y finales) y descarta el cuerpo. Esto permite subir el
       libro entero sin recortar partes a mano (ver app/zonas.py).
    5. Convierte a imágenes las páginas de esas zonas, siempre que se pueda,
       de modo que la portada y el colofón lleguen al modelo de visión aunque
       el libro tenga cientos de páginas.
    6. Devuelve una Entrada lista para el extractor, con cobertura registrada.

Este módulo es la primera línea de defensa. Debe ser estricto: cualquier
formato ambiguo se rechaza en lugar de enviarlo al parser documental.
"""

from __future__ import annotations

import base64
import io
import logging
import os
import re
import warnings
import zipfile
from dataclasses import dataclass, field
from typing import Literal

from . import zonas
from .extractor import Entrada
from .parser_sandbox import SandboxExecutionError, ejecutar_en_sandbox, sandbox_activo

logger = logging.getLogger(__name__)


# =============================================================================
# Límites de seguridad
# =============================================================================

TAMANO_MAXIMO_BYTES = int(os.getenv("MAX_TAMANO_FICHERO_BYTES", str(80 * 1024 * 1024)))

# Con segmentación por zonas ya no se procesa el volumen completo: el texto se
# recorta a las zonas y solo se rasterizan ~14 páginas. Por eso el tope de
# páginas puede ser alto: el coste real depende de las zonas, no del total.
PAGINAS_MAXIMAS_PDF = int(os.getenv("PAGINAS_MAXIMAS_PDF", "1200"))
LONGITUD_MAXIMA_TEXTO = int(os.getenv("MAX_LONGITUD_TEXTO_EXTRAIDO", "800000"))

# Umbral de calidad OCR por debajo del cual se usa ruta visión en vez de texto.
UMBRAL_CALIDAD_OCR = float(os.getenv("UMBRAL_CALIDAD_OCR", "0.5"))

# Imagen de entrada. 40 MP evita bombas razonables sin impedir escaneos
# administrativos grandes. La dimensión máxima reduce cargas patológicas.
DIMENSION_MAXIMA_IMAGEN = 6_000
PIXELS_MAXIMOS_IMAGEN = 40_000_000
BYTES_MAXIMOS_IMAGEN_NORMALIZADA = 12 * 1024 * 1024

# Render PDF. Se calculan píxeles estimados antes de renderizar.
ESCALA_RENDER_PDF = 2.0
PIXELS_MAXIMOS_PAGINA_PDF = 35_000_000
BYTES_MAXIMOS_IMAGEN_PDF = 12 * 1024 * 1024
BYTES_MAXIMOS_TOTAL_IMAGENES = 120 * 1024 * 1024

# DOCX/ZIP. Límites deliberadamente conservadores para evitar zip bombs.
DOCX_MAX_ENTRADAS = 500
DOCX_MAX_DESCOMPRIMIDO = 30 * 1024 * 1024
DOCX_MAX_RATIO_COMPRESION = 100
DOCX_MAX_DOCUMENT_XML = 5 * 1024 * 1024
DOCX_MAX_MEDIA_TOTAL = 20 * 1024 * 1024

TIPOS_ADMITIDOS = {
    "application/pdf":              "pdf",
    "text/plain":                   "texto",
    "text/markdown":                "texto",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/jpeg":                   "imagen",
    "image/png":                    "imagen",
    "image/tiff":                   "imagen",
    "image/webp":                   "imagen",
}

Ruta = Literal["texto", "vision", "hibrida"]


# =============================================================================
# Excepciones
# =============================================================================

class ErrorValidacion(Exception):
    """Se lanza cuando el fichero no supera las comprobaciones de seguridad."""


# =============================================================================
# Resultado
# =============================================================================

@dataclass
class DocumentoProcesado:
    entrada: Entrada
    ruta: Ruta
    nombre_original: str
    tipo_mime: str
    tamano_bytes: int
    paginas: int | None
    # Cobertura: qué se analizó realmente frente al documento completo.
    # Lo consume la ficha técnica de auditoría.
    cobertura: dict = field(default_factory=dict)


# =============================================================================
# Validación
# =============================================================================

def validar(contenido: bytes, nombre: str) -> str:
    """
    Valida el fichero y devuelve su tipo MIME real. No se acepta fallback
    por extensión: si la firma/contenido no encaja, se rechaza.
    """
    if len(contenido) == 0:
        raise ErrorValidacion("El fichero está vacío.")
    if len(contenido) > TAMANO_MAXIMO_BYTES:
        raise ErrorValidacion(
            f"El fichero supera el tamaño máximo "
            f"({len(contenido) / 1_048_576:.1f} MB, máximo "
            f"{TAMANO_MAXIMO_BYTES / 1_048_576:.0f} MB)."
        )

    mime = _detectar_mime_real(contenido)

    if mime not in TIPOS_ADMITIDOS:
        raise ErrorValidacion(
            f"Tipo de fichero no admitido: {mime}. "
            f"Formatos soportados: PDF, DOCX, TXT, JPG, PNG, TIFF, WebP."
        )

    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        _validar_docx_seguro(contenido)

    return mime


def _detectar_mime_real(contenido: bytes) -> str:
    if len(contenido) < 4:
        raise ErrorValidacion("Fichero demasiado pequeño para ser válido.")

    if contenido.startswith(b"%PDF-"):
        return "application/pdf"
    if contenido.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if contenido.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if contenido.startswith((b"II*\x00", b"MM\x00*")):
        return "image/tiff"
    if contenido.startswith(b"RIFF") and len(contenido) >= 12 and contenido[8:12] == b"WEBP":
        return "image/webp"
    if contenido.startswith(b"PK\x03\x04"):
        if _es_docx(contenido):
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        raise ErrorValidacion("Fichero ZIP no reconocido como DOCX admitido.")

    if _parece_texto(contenido):
        return "text/plain"

    raise ErrorValidacion(
        "No se ha podido determinar el tipo del fichero. "
        "Asegúrese de que es un PDF, DOCX, TXT o imagen válida."
    )


def _es_docx(contenido: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            nombres = set(z.namelist())
    except zipfile.BadZipFile:
        return False
    return "[Content_Types].xml" in nombres and "word/document.xml" in nombres


def _validar_docx_seguro(contenido: bytes) -> None:
    """Valida estructura DOCX antes de entregarla a python-docx."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        raise ErrorValidacion("El documento DOCX está dañado o no es válido.") from None

    if not infos:
        raise ErrorValidacion("El DOCX no contiene entradas internas válidas.")
    if len(infos) > DOCX_MAX_ENTRADAS:
        raise ErrorValidacion(
            f"El DOCX contiene demasiadas entradas internas ({len(infos)}; "
            f"máximo {DOCX_MAX_ENTRADAS})."
        )

    total_descomprimido = 0
    total_media = 0
    tiene_document_xml = False

    for info in infos:
        nombre = info.filename.replace("\\", "/")

        if nombre.startswith("/") or ".." in nombre.split("/"):
            raise ErrorValidacion("El DOCX contiene rutas internas no seguras.")

        if nombre.endswith("/"):
            continue

        if "vbaProject" in nombre:
            raise ErrorValidacion(
                "El documento contiene macros. Por seguridad, no se procesan "
                "documentos con macros. Guárdelo como .docx sin macros."
            )

        total_descomprimido += info.file_size
        if total_descomprimido > DOCX_MAX_DESCOMPRIMIDO:
            raise ErrorValidacion(
                "El DOCX se expande demasiado al descomprimirse. "
                "Puede estar dañado o construido de forma maliciosa."
            )

        if info.compress_size > 0:
            ratio = info.file_size / info.compress_size
            if ratio > DOCX_MAX_RATIO_COMPRESION and info.file_size > 1_000_000:
                raise ErrorValidacion(
                    "El DOCX tiene una ratio de compresión anómala. "
                    "Puede tratarse de un ZIP-bomb."
                )

        if nombre == "word/document.xml":
            tiene_document_xml = True
            if info.file_size > DOCX_MAX_DOCUMENT_XML:
                raise ErrorValidacion("El cuerpo XML principal del DOCX es demasiado grande.")

        if nombre.startswith("word/media/"):
            total_media += info.file_size
            if total_media > DOCX_MAX_MEDIA_TOTAL:
                raise ErrorValidacion("El DOCX contiene demasiados recursos multimedia.")

    if not tiene_document_xml:
        raise ErrorValidacion("El DOCX no contiene word/document.xml.")


def _parece_texto(contenido: bytes) -> bool:
    """
    Detección estricta de texto. Evita aceptar binarios aleatorios por el
    simple hecho de que Latin-1 pueda decodificar cualquier byte.
    """
    muestra = contenido[:8192]
    if not muestra:
        return False

    if muestra.startswith((b"\xff\xfe", b"\xfe\xff")):
        return False

    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            texto = muestra.decode(encoding)
        except UnicodeDecodeError:
            continue

        if "\x00" in texto:
            return False

        controles = sum(1 for c in texto if ord(c) < 32 and c not in "\n\r\t")
        if controles > max(2, len(texto) * 0.01):
            return False

        imprimibles = sum(1 for c in texto if c.isprintable() or c in "\n\r\t")
        if imprimibles / max(1, len(texto)) < 0.95:
            return False

        letras_o_numeros = sum(1 for c in texto if c.isalnum())
        espacios = sum(1 for c in texto if c.isspace())
        if letras_o_numeros < max(8, len(texto) * 0.20):
            return False
        if len(texto) > 200 and espacios == 0:
            return False

        return True

    return False


# =============================================================================
# Evaluación de calidad del OCR
# =============================================================================

def _calidad_ocr(texto: str, num_paginas: int) -> float:
    """Estima la calidad del texto OCR en una escala 0.0-1.0."""
    if not texto or num_paginas == 0:
        return 0.0

    chars_por_pagina = len(texto) / num_paginas
    if chars_por_pagina < 200:
        return 0.0
    densidad = min(1.0, chars_por_pagina / 600)

    alfanum = sum(1 for c in texto if c.isalnum() or c in " .,;:áéíóúñüÁÉÍÓÚÑÜ¿?¡!")
    proporcion_sana = alfanum / max(1, len(texto))

    palabras = [p for p in texto.split() if len(p) > 3 and any(c.isalpha() for c in p)]
    palabras_por_pagina = len(palabras) / num_paginas
    palabras_score = min(1.0, palabras_por_pagina / 80)

    lineas = [linea.strip() for linea in texto.split("\n") if linea.strip()]
    if lineas:
        cortas = sum(1 for linea in lineas if len(linea.split()) <= 2)
        prop_cortas = cortas / len(lineas)
        lineas_score = max(0.0, 1.0 - max(0.0, prop_cortas - 0.4) * 2)
    else:
        lineas_score = 0.0

    return (densidad * 0.3 + proporcion_sana * 0.3
            + palabras_score * 0.25 + lineas_score * 0.15)


# =============================================================================
# Extracción por tipo
# =============================================================================

def _extraer_texto_pdf_por_pagina(contenido: bytes) -> list[str]:
    """
    Extrae el texto de un PDF página a página. Devuelve una lista paralela al
    índice de página; esto es lo que necesita la segmentación por zonas para
    saber qué texto pertenece a preliminares y qué a finales.
    """
    import pypdf

    try:
        lector = pypdf.PdfReader(io.BytesIO(contenido), strict=False)
    except Exception as e:
        raise ErrorValidacion(f"El PDF no se puede abrir: {e}") from None

    if lector.is_encrypted:
        raise ErrorValidacion("El PDF está cifrado o protegido por contraseña.")

    num_paginas = len(lector.pages)
    if num_paginas <= 0:
        raise ErrorValidacion("El PDF no contiene páginas.")
    if num_paginas > PAGINAS_MAXIMAS_PDF:
        raise ErrorValidacion(
            f"El PDF tiene {num_paginas} páginas; el máximo admitido es "
            f"{PAGINAS_MAXIMAS_PDF}. Divida el documento en piezas más pequeñas."
        )

    paginas: list[str] = []
    for pagina in lector.pages:
        try:
            paginas.append(pagina.extract_text() or "")
        except Exception as e:
            logger.warning("Error extrayendo texto de una página: %s", e)
            paginas.append("")
    return paginas


def _pdf_paginas_a_imagenes(contenido: bytes, indices: list[int]) -> list[bytes]:
    """
    Rasteriza únicamente las páginas indicadas (por índice 0-based). A
    diferencia de un render secuencial, esto permite incluir el colofón
    (al final del volumen) sin renderizar todo lo que hay en medio.
    """
    import pypdfium2 as pdfium

    imagenes: list[bytes] = []
    total_bytes = 0

    try:
        pdf = pdfium.PdfDocument(contenido)
    except Exception as e:
        raise ErrorValidacion(f"El PDF no se puede abrir para renderizado: {e}") from None

    try:
        n = len(pdf)
        for i in indices:
            if i < 0 or i >= n:
                continue
            pagina = pdf[i]
            try:
                ancho_pt, alto_pt = pagina.get_size()
            except Exception:
                ancho_pt, alto_pt = (0, 0)

            if ancho_pt and alto_pt:
                pixeles_estimados = (
                    int(ancho_pt * ESCALA_RENDER_PDF) * int(alto_pt * ESCALA_RENDER_PDF)
                )
                if pixeles_estimados > PIXELS_MAXIMOS_PAGINA_PDF:
                    raise ErrorValidacion(
                        f"La página {i + 1} del PDF es demasiado grande para renderizar "
                        f"({pixeles_estimados:,} píxeles estimados)."
                    )

            bitmap = pagina.render(scale=ESCALA_RENDER_PDF).to_pil()
            try:
                if bitmap.width * bitmap.height > PIXELS_MAXIMOS_PAGINA_PDF:
                    raise ErrorValidacion(
                        f"La página {i + 1} del PDF supera el límite de píxeles renderizados."
                    )
                if bitmap.mode not in ("RGB", "L"):
                    bitmap = bitmap.convert("RGB")
                buf = io.BytesIO()
                bitmap.save(buf, format="PNG", optimize=True)
                datos = buf.getvalue()
            finally:
                try:
                    bitmap.close()
                except Exception:
                    pass

            if len(datos) > BYTES_MAXIMOS_IMAGEN_PDF:
                raise ErrorValidacion(
                    f"La imagen renderizada de la página {i + 1} es demasiado grande."
                )
            total_bytes += len(datos)
            if total_bytes > BYTES_MAXIMOS_TOTAL_IMAGENES:
                raise ErrorValidacion(
                    "El conjunto de imágenes renderizadas del PDF es demasiado grande."
                )
            imagenes.append(datos)
    finally:
        pdf.close()

    return imagenes


def _extraer_texto_docx(contenido: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(contenido))
    except Exception as e:
        raise ErrorValidacion(f"El DOCX no se puede abrir: {e}") from None

    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))

    return "\n".join(partes).strip()


def _validar_imagen(contenido: bytes) -> bytes:
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = PIXELS_MAXIMOS_IMAGEN

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(contenido)) as img_probe:
                img_probe.verify()

            with Image.open(io.BytesIO(contenido)) as img:
                img.load()

                ancho, alto = img.size
                pixeles = ancho * alto
                if pixeles > PIXELS_MAXIMOS_IMAGEN:
                    raise ErrorValidacion(
                        f"La imagen es demasiado grande ({pixeles:,} píxeles; "
                        f"máximo {PIXELS_MAXIMOS_IMAGEN:,})."
                    )

                if getattr(img, "n_frames", 1) > 1:
                    if img.format in {"TIFF", "WEBP"}:
                        raise ErrorValidacion("No se admiten imágenes multipágina o animadas.")

                procesada = img
                creada = None
                try:
                    if max(ancho, alto) > DIMENSION_MAXIMA_IMAGEN:
                        factor = DIMENSION_MAXIMA_IMAGEN / max(ancho, alto)
                        nuevo = (max(1, int(ancho * factor)), max(1, int(alto * factor)))
                        creada = procesada.resize(nuevo, Image.Resampling.LANCZOS)
                        procesada = creada

                    if procesada.mode not in ("RGB", "L"):
                        convertida = procesada.convert("RGB")
                        if creada is not None:
                            creada.close()
                        creada = convertida
                        procesada = convertida

                    buf = io.BytesIO()
                    procesada.save(buf, format="PNG", optimize=True)
                    normalizada = buf.getvalue()
                finally:
                    if creada is not None:
                        try:
                            creada.close()
                        except Exception:
                            pass

                if len(normalizada) > BYTES_MAXIMOS_IMAGEN_NORMALIZADA:
                    raise ErrorValidacion("La imagen normalizada es demasiado grande para procesarse.")
                return normalizada
    except ErrorValidacion:
        raise
    except Exception as e:
        raise ErrorValidacion(f"La imagen no es válida o está dañada: {e}") from None


def _limpiar_texto(texto: str) -> str:
    if len(texto) > LONGITUD_MAXIMA_TEXTO:
        logger.info("Texto truncado de %d a %d caracteres", len(texto), LONGITUD_MAXIMA_TEXTO)
        texto = texto[:LONGITUD_MAXIMA_TEXTO]

    limpio = "".join(c for c in texto if c.isprintable() or c in "\n\t\r")
    limpio = re.sub(r"\n{4,}", "\n\n\n", limpio)
    return limpio.strip()


# =============================================================================
# Serialización segura del resultado del sandbox
# =============================================================================

def _documento_a_payload(doc: DocumentoProcesado) -> dict:
    return {
        "entrada": {
            "texto": doc.entrada.texto,
            "imagenes": [base64.b64encode(img).decode("ascii") for img in (doc.entrada.imagenes or [])],
            "imagenes_etiquetas": doc.entrada.imagenes_etiquetas,
            "plantilla": doc.entrada.plantilla,
            "instrucciones_tipo": doc.entrada.instrucciones_tipo,
        },
        "ruta": doc.ruta,
        "nombre_original": doc.nombre_original,
        "tipo_mime": doc.tipo_mime,
        "tamano_bytes": doc.tamano_bytes,
        "paginas": doc.paginas,
        "cobertura": doc.cobertura,
    }


def _documento_desde_payload(payload: dict) -> DocumentoProcesado:
    entrada_payload = payload.get("entrada", {})
    imagenes_b64 = entrada_payload.get("imagenes") or []
    imagenes = [base64.b64decode(img) for img in imagenes_b64]
    return DocumentoProcesado(
        entrada=Entrada(
            texto=entrada_payload.get("texto"),
            imagenes=imagenes or None,
            imagenes_etiquetas=entrada_payload.get("imagenes_etiquetas") or None,
            plantilla=entrada_payload.get("plantilla"),
            instrucciones_tipo=entrada_payload.get("instrucciones_tipo") or {},
        ),
        ruta=payload["ruta"],
        nombre_original=payload["nombre_original"],
        tipo_mime=payload["tipo_mime"],
        tamano_bytes=int(payload["tamano_bytes"]),
        paginas=payload.get("paginas"),
        cobertura=payload.get("cobertura") or {},
    )


def _procesar_impl_serializable(contenido: bytes, nombre: str) -> dict:
    """Versión serializable para el sandbox: no devuelve objetos Python."""
    return _documento_a_payload(_procesar_impl(contenido, nombre))


# =============================================================================
# Función pública
# =============================================================================

def _procesar_pdf(contenido: bytes) -> tuple[Entrada, Ruta, int, dict]:
    """
    Procesa un PDF completo aplicando segmentación por zonas.

    Devuelve (entrada, ruta, num_paginas, cobertura).
    """
    paginas_texto = _extraer_texto_pdf_por_pagina(contenido)
    num_paginas = len(paginas_texto)

    zonif = zonas.segmentar_pdf(paginas_texto)
    texto_zonificado_crudo = zonas.construir_texto_zonificado(zonif)
    calidad = _calidad_ocr(texto_zonificado_crudo, max(1, zonif.paginas_analizadas))
    logger.info(
        "PDF: %d págs totales, %d analizadas por zonas, calidad OCR zonas %.2f",
        num_paginas, zonif.paginas_analizadas, calidad,
    )

    cobertura = zonif.resumen()
    cobertura["calidad_ocr_zonas"] = round(calidad, 3)
    cobertura["umbral_calidad_ocr"] = UMBRAL_CALIDAD_OCR

    indices_vision = zonas.indices_para_vision(zonif)
    imagenes: list[bytes] | None = None
    etiquetas: list[str] | None = None
    try:
        imagenes = _pdf_paginas_a_imagenes(contenido, indices_vision)
        etiquetas = zonas.etiquetas_para_indices(zonif, indices_vision)
    except ErrorValidacion:
        raise
    except Exception as e:
        logger.warning("No se pudieron rasterizar las zonas del PDF: %s", e)
        imagenes = None
        etiquetas = None

    texto_util = calidad >= UMBRAL_CALIDAD_OCR and bool(texto_zonificado_crudo)

    if texto_util and imagenes:
        ruta: Ruta = "hibrida"
        texto = _limpiar_texto(texto_zonificado_crudo)
    elif texto_util:
        ruta = "texto"
        texto = _limpiar_texto(texto_zonificado_crudo)
    elif imagenes:
        ruta = "vision"
        texto = None
        logger.info("Texto OCR de las zonas insuficiente; ruta visión sobre zonas")
    else:
        raise ErrorValidacion(
            "El PDF no tiene texto legible en sus zonas descriptivas y no se "
            "pudo convertir a imagen. Puede estar dañado o ser demasiado complejo."
        )

    cobertura["ruta"] = ruta
    cobertura["paginas_rasterizadas"] = len(imagenes) if imagenes else 0

    entrada = Entrada(texto=texto, imagenes=imagenes, imagenes_etiquetas=etiquetas)
    return entrada, ruta, num_paginas, cobertura


def _procesar_impl(contenido: bytes, nombre: str) -> DocumentoProcesado:
    """Implementación real del procesamiento documental.

    Esta función se ejecuta normalmente dentro de un proceso aislado.
    """
    mime = validar(contenido, nombre)
    familia = TIPOS_ADMITIDOS[mime]

    texto: str | None = None
    imagenes: list[bytes] | None = None
    etiquetas: list[str] | None = None
    ruta: Ruta
    paginas: int | None = None
    cobertura: dict = {}

    if familia == "pdf":
        entrada, ruta, paginas, cobertura = _procesar_pdf(contenido)
        return DocumentoProcesado(
            entrada=entrada, ruta=ruta, nombre_original=nombre,
            tipo_mime=mime, tamano_bytes=len(contenido), paginas=paginas,
            cobertura=cobertura,
        )

    elif familia == "docx":
        texto_completo = _limpiar_texto(_extraer_texto_docx(contenido))
        if not texto_completo:
            raise ErrorValidacion("El documento DOCX no contiene texto legible.")
        texto, cobertura = zonas.segmentar_texto_plano(texto_completo)
        cobertura["ruta"] = "texto"
        ruta = "texto"

    elif familia == "texto":
        texto_completo = _limpiar_texto(_decodificar_texto(contenido))
        if not texto_completo:
            raise ErrorValidacion("El fichero de texto está vacío.")
        texto, cobertura = zonas.segmentar_texto_plano(texto_completo)
        cobertura["ruta"] = "texto"
        ruta = "texto"

    elif familia == "imagen":
        # Una imagen suelta ya es una parte concreta del libro: no se segmenta.
        imagenes = [_validar_imagen(contenido)]
        ruta = "vision"
        cobertura = {"estrategia": "imagen_suelta", "ruta": "vision",
                     "paginas_totales": 1, "paginas_analizadas": 1}

    else:
        raise ErrorValidacion(f"Familia de tipo no soportado: {familia}")

    entrada = Entrada(texto=texto, imagenes=imagenes, imagenes_etiquetas=etiquetas)

    return DocumentoProcesado(
        entrada=entrada,
        ruta=ruta,
        nombre_original=nombre,
        tipo_mime=mime,
        tamano_bytes=len(contenido),
        paginas=paginas,
        cobertura=cobertura,
    )


def _decodificar_texto(contenido: bytes) -> str:
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return contenido.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ErrorValidacion("No se pudo decodificar el fichero de texto.")


def procesar(contenido: bytes, nombre: str) -> DocumentoProcesado:
    """
    Valida y prepara el fichero para el extractor.

    Por defecto delega en un proceso hijo aislado. Esto reduce el impacto de
    bloqueos, fugas nativas o consumos anómalos en parsers de PDF/DOCX/imagen.
    Puede desactivarse con USAR_SANDBOX_PARSERS=false únicamente para depuración.
    """
    if (
        sandbox_activo()
        and os.getenv("_TIPO_SANDBOX_CHILD") != "1"
    ):
        try:
            payload = ejecutar_en_sandbox("app.router:_procesar_impl_serializable", contenido, nombre)
            if not isinstance(payload, dict):
                raise ErrorValidacion("El parser aislado devolvió una respuesta inválida.")
            return _documento_desde_payload(payload)
        except SandboxExecutionError as exc:
            if exc.exception_type == "ErrorValidacion":
                raise ErrorValidacion(exc.message) from None
            logger.warning(
                "Parser aislado falló (%s): %s",
                exc.exception_type,
                exc.message,
            )
            raise ErrorValidacion(
                "No se pudo procesar el documento dentro de los límites de seguridad. "
                "Puede estar dañado, ser demasiado complejo o requerir división previa."
            ) from None

    return _procesar_impl(contenido, nombre)
